"""Extract ink workbook to a local web data module. Never modify the workbook."""
import hashlib
import json
import sys
from pathlib import Path
import openpyxl

source, output = map(Path, sys.argv[1:3])
book = openpyxl.load_workbook(source, data_only=True)

def values(sheet, row, count):
    return [sheet.cell(row, col).value for col in range(1, count + 1)]

purchase = book['รายการซื้อ']
groups = []
for year, title_row, note_row, first, last, total_row in [(2568, 2, 3, 5, 12, 13), (2569, 15, 16, 18, 29, 30), (2569, 31, 32, 34, 45, 46)]:
    rows = [{'sourceRow': i, 'cells': values(purchase, i, 5)} for i in range(first, last + 1)]
    total = purchase.cell(total_row, 5).value
    assert round(sum(row['cells'][4] for row in rows), 2) == total
    for row in rows:
        assert round(row['cells'][2] * row['cells'][3], 2) == row['cells'][4]
    groups.append({'id': str(title_row), 'year': year, 'title': purchase.cell(title_row, 1).value,
                   'note': purchase.cell(note_row, 1).value, 'rows': rows, 'total': total})

issues = book['รายการเบิกหมึก']
def merged_value(row, col):
    value = issues.cell(row, col).value
    if value is not None:
        return value
    for span in issues.merged_cells.ranges:
        if span.min_row <= row <= span.max_row and span.min_col <= col <= span.max_col:
            return issues.cell(span.min_row, span.min_col).value
    return None

withdrawals = []
for i in range(3, issues.max_row + 1):
    if issues.cell(i, 3).value is None:
        continue
    date, department = merged_value(i, 1), merged_value(i, 2)
    assert date and department
    assert date.year in (2568, 2569), 'Expected source Buddhist dates'
    withdrawals.append({'sourceRow': i, 'dateISO': f'{date.year - 543:04d}-{date.month:02d}-{date.day:02d}',
                        'year': date.year + (date.month >= 10),
                        'cells': [f'{date.day:02d}/{date.month:02d}/{date.year}', department, issues.cell(i, 3).value, issues.cell(i, 4).value]})

stock = book['คงเหลือ']
balances = [{'sourceRow': i, 'cells': values(stock, i, 5)} for i in range(3, 17)]
for row in balances:
    assert round(row['cells'][2] * row['cells'][3], 2) == row['cells'][4]
assert sum(row['cells'][4] for row in balances) == stock['E17'].value == 137160
assert len(withdrawals) == 36
assert sum(row['cells'][3] for row in withdrawals) == 35
data = {'sourceName': source.name, 'sourceSha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'purchaseTitle': purchase['A1'].value, 'purchaseColumns': values(purchase, 4, 5), 'groups': groups,
        'purchaseNotes': [purchase.cell(i, 2).value for i in (48, 49, 50)],
        'withdrawalTitle': issues['A1'].value, 'withdrawalColumns': values(issues, 2, 4), 'withdrawals': withdrawals,
        'stockTitle': stock['A1'].value, 'stockColumns': values(stock, 2, 5), 'balances': balances,
        'stockTotalLabel': stock['D17'].value, 'stockTotal': stock['E17'].value}
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text('// Generated from the supplied workbook; source remains unchanged.\nexport const inkData = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
print(json.dumps({'purchaseRows': 32, 'withdrawalRows': 36, 'stockRows': 14, 'purchaseTotals': [g['total'] for g in groups], 'stockValue': data['stockTotal']}))
