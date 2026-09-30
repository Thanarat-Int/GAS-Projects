from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path

import openpyxl


CATEGORIES = {
    "สนง.": ("CAT-OFFICE", "OFF", "วัสดุสำนักงาน", "Office supplies", 10),
    "การแพทย์": ("CAT-MEDICAL", "MED", "วัสดุการแพทย์", "Medical supplies", 20),
    "งานบ้าน": ("CAT-HOUSEKEEPING", "HSK", "วัสดุงานบ้าน", "Housekeeping supplies", 30),
    "สนง.พิมพ์": ("CAT-PRINTED", "PRT", "วัสดุสำนักงาน (แบบพิมพ์)", "Printed office forms", 40),
    "คอม": ("CAT-IT", "IT", "วัสดุคอมพิวเตอร์", "Computer supplies", 50),
    "เวชภัณฑ์ (ชั้น5)": ("CAT-MEDSUP-5", "MS5", "เวชภัณฑ์ (ชั้น 5)", "Medical supplies (5th floor)", 60),
    "เวชภัณฑ์": ("CAT-MEDSUP", "MSP", "เวชภัณฑ์ (มิใช่ยา)", "Non-drug medical supplies", 70),
}

UNIT_MAP = {
    "BOTTLE": "ขวด",
    "BOX": "กล่อง",
    "PIECE": "ชิ้น",
    "CAN": "กระป๋อง",
    "PACK": "แพ็ค",
    "GALLON": "แกลลอน",
    "BAG": "ถุง",
}


def numeric(value: object, *, dash_is_zero: bool = True) -> tuple[float, bool]:
    if value in (None, ""):
        return 0.0, True
    if dash_is_zero and isinstance(value, str) and value.strip() == "-":
        return 0.0, False
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return float(value), True
    if isinstance(value, str):
        try:
            return float(value.replace(",", "").strip()), False
        except ValueError:
            pass
    raise ValueError(f"Unsupported numeric value: {value!r}")


def js_value(value: object) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(int(value)) if float(value).is_integer() else repr(float(value))
    return json.dumps(str(value), ensure_ascii=False)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    source = Path(args.input)
    target = Path(args.output)
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=False)
    rows: list[list[object]] = []

    for sheet_name, category in CATEGORIES.items():
        sheet = workbook[sheet_name]
        for row_number in range(1, sheet.max_row + 1):
            sequence = sheet.cell(row_number, 1).value
            source_name = sheet.cell(row_number, 2).value
            source_unit = sheet.cell(row_number, 3).value
            if not isinstance(sequence, (int, float)) or source_name in (None, "") or source_unit in (None, ""):
                continue

            source_name = str(source_name).strip()
            source_unit = str(source_unit).strip()
            proposed_unit = UNIT_MAP.get(source_unit.upper(), source_unit)
            opening, opening_clean = numeric(sheet.cell(row_number, 4).value)
            received, received_clean = numeric(sheet.cell(row_number, 5).value)
            source_total, total_clean = numeric(sheet.cell(row_number, 7).value)
            issued, issued_clean = numeric(sheet.cell(row_number, 9).value)
            balance, balance_clean = numeric(sheet.cell(row_number, 10).value)
            price, price_clean = numeric(sheet.cell(row_number, 12).value)
            source_value, value_clean = numeric(sheet.cell(row_number, 14).value)

            issues: list[str] = []
            if not opening_clean:
                issues.append("OPENING_TEXT")
            if not all((received_clean, total_clean, issued_clean, balance_clean, price_clean, value_clean)):
                issues.append("TEXT_CONVERTED")
            if abs(source_total - (opening + received)) > 1e-8:
                issues.append("SOURCE_TOTAL_CONFLICT")
            if abs(balance - (opening + received - issued)) > 1e-8:
                issues.append("SOURCE_BALANCE_CONFLICT")
            if price <= 0:
                issues.append("MISSING_PRICE")
            elif abs(source_value - balance * price) > 0.011:
                issues.append("SOURCE_VALUE_CONFLICT")
            if proposed_unit != source_unit:
                issues.append("UNIT_MAPPED")
            if balance < 0:
                issues.append("NEGATIVE_BALANCE")

            review_id = f"IMP-{len(rows) + 1:04d}"
            rows.append([
                review_id,
                f"{sheet_name}:{row_number}",
                sheet_name,
                row_number,
                category[0],
                source_name,
                source_unit,
                source_name,
                proposed_unit,
                opening,
                received,
                source_total,
                issued,
                balance,
                price,
                source_value,
                "|".join(issues),
                "PENDING",
                "",
                "",
                "",
                "",
                "",
                1,
            ])

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    lines = [
        "/** Generated from the supplied annual material workbook. Do not edit manually. */",
        "var MATERIAL_SEED_META = Object.freeze({",
        f"  sourceName: {js_value(source.name)},",
        f"  sourceSha256: {js_value(digest)},",
        f"  recordCount: {len(rows)}",
        "});",
        "",
        "var MATERIAL_CATEGORY_SEED = Object.freeze([",
    ]
    lines.extend("  [" + ", ".join(js_value(value) for value in row) + "]," for row in CATEGORIES.values())
    lines.extend(["]);", "", "var MATERIAL_IMPORT_SEED = Object.freeze(["])
    lines.extend("  [" + ", ".join(js_value(value) for value in row) + "]," for row in rows)
    lines.extend(["]);", ""])
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"output": str(target), "records": len(rows), "sha256": digest}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
