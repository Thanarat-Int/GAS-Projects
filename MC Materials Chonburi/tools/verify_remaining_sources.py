"""Read all five requested sheets without changing the workbook or seed."""
import json
import sys
from pathlib import Path
import openpyxl
from generate_material_seed import numeric, UNIT_MAP, CATEGORIES

sys.stdout.reconfigure(encoding='utf-8')
workbook = openpyxl.load_workbook(sys.argv[1], data_only=True)
seed = [json.loads(line.strip().rstrip(',')) for line in Path(sys.argv[2]).read_text(encoding='utf-8').splitlines()
        if line.strip().startswith('["IMP-')]
counts = {'งานบ้าน': 28, 'สนง.พิมพ์': 4, 'คอม': 1, 'เวชภัณฑ์ (ชั้น5)': 6, 'เวชภัณฑ์': 5}
for name, count in counts.items():
    sheet = workbook[name]
    rows = [r for r in seed if r[4] == CATEGORIES[name][0]]
    indices = [i for i in range(1, sheet.max_row + 1)
               if isinstance(sheet.cell(i, 1).value, (int, float)) and sheet.cell(i, 2).value]
    assert len(rows) == len(indices) == count, name
    assert [r[3] for r in rows] == indices, name
    for row in rows:
        i = row[3]
        material, unit = [str(sheet.cell(i, c).value).strip() for c in (2, 3)]
        assert row[5:9] == [material, unit, material, UNIT_MAP.get(unit.upper(), unit)], (name, i)
        assert row[9:16] == [numeric(sheet.cell(i, c).value)[0] for c in (4, 5, 7, 9, 10, 12, 14)], (name, i)
        assert not sheet.cell(i, 15).value, (name, i, 'Unmapped note')
    total = round(sum(r[15] for r in rows), 2)
    source_total = sheet.cell(indices[-1] + 1, 14).value
    if source_total is not None:
        assert total == source_total, name
    print(json.dumps({'sheet': name, 'rows': count, 'value': total, 'sourceTotal': source_total,
                      'dateHeading': sheet['A4'].value, 'fiscalHeading': sheet['J6'].value,
                      'conflicts': [{'row': r[3], 'issues': r[16]} for r in rows if 'CONFLICT' in r[16]],
                      'missingPrices': [r[3] for r in rows if sheet.cell(r[3], 12).value is None]}, ensure_ascii=False))
