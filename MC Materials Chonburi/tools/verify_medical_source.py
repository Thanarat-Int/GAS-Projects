"""Read-only reconciliation of the medical sheet against the checked-in seed."""
import json
import sys
from pathlib import Path
import openpyxl
from generate_material_seed import numeric, UNIT_MAP

sys.stdout.reconfigure(encoding='utf-8')
workbook = openpyxl.load_workbook(sys.argv[1], data_only=True)
sheet = workbook['การแพทย์']
seed = [json.loads(line.strip().rstrip(',')) for line in Path(sys.argv[2]).read_text(encoding='utf-8').splitlines()
        if line.strip().startswith('["IMP-')]
medical = [row for row in seed if row[4] == 'CAT-MEDICAL']
source_rows = [i for i in range(1, sheet.max_row + 1)
               if isinstance(sheet.cell(i, 1).value, (int, float)) and sheet.cell(i, 2).value]
assert len(medical) == len(source_rows) == 28
assert [row[3] for row in medical] == source_rows
for row in medical:
    index = row[3]
    name, unit = [str(sheet.cell(index, c).value).strip() for c in (2, 3)]
    assert row[5:9] == [name, unit, name, UNIT_MAP.get(unit.upper(), unit)], index
    expected = [numeric(sheet.cell(index, c).value)[0] for c in (4, 5, 7, 9, 10, 12, 14)]
    assert row[9:16] == expected, index
    assert not sheet.cell(index, 15).value, 'Unmapped source note'
assert round(sum(row[15] for row in medical), 2) == sheet['N38'].value == 9052.1
print(json.dumps({'verifiedRows': len(medical), 'sourceRange': 'การแพทย์!A10:O37',
                  'reportedValue': 9052.1, 'zeroPriceItems': [r[5] for r in medical if not r[14]],
                  'duplicateSequence': 'Source sequence 8 occurs twice; both rows retained.'}, ensure_ascii=False))
